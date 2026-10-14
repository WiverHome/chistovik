"""Генератор шаблона курсовой / ВКР по правилам оформления из JSON-конфига.

Использование:
    python make_template.py config/kursovaya.json -o out/Курсовая.docx
    python make_template.py config/vkr.json -o out/ВКР.docx --pdf

Если установлен Microsoft Word, после сборки он обновляет оглавление и номера
рисунков/таблиц (и при --pdf сохраняет PDF). Без Word поля обновятся при открытии файла.
"""
import argparse
import datetime
import json
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm, Mm, Pt, RGBColor

ALIGN = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
}
TWIPS_PER_CM = 566.93

# Порядок дочерних элементов w:pPr по схеме OOXML (Word строго его проверяет)
PPR_ORDER = [
    "w:pStyle", "w:keepNext", "w:keepLines", "w:pageBreakBefore", "w:framePr", "w:widowControl",
    "w:numPr", "w:suppressLineNumbers", "w:pBdr", "w:shd", "w:tabs", "w:suppressAutoHyphens",
    "w:kinsoku", "w:wordWrap", "w:overflowPunct", "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN",
    "w:bidi", "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind", "w:contextualSpacing",
    "w:mirrorIndents", "w:suppressOverlap", "w:jc", "w:textDirection", "w:textAlignment",
    "w:textboxTightWrap", "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange",
]


# ---------------------------------------------------------------- конфиг

def deep_merge(base, over):
    out = dict(base)
    for k, v in over.items():
        out[k] = deep_merge(out[k], v) if isinstance(v, dict) and isinstance(out.get(k), dict) else v
    return out


def load_config(path: Path) -> dict:
    data = json.loads(path.read_text(encoding="utf-8"))
    if "extends" in data:
        data = deep_merge(load_config(path.parent / data.pop("extends")), data)
    return data


class Fields(dict):
    def __missing__(self, key):
        return "{" + key + "}"


# ---------------------------------------------------------------- низкоуровневые помощники

def ppr_set(pPr, tag, attrs=None):
    el = pPr.find(qn(tag))
    if el is None:
        el = OxmlElement(tag)
        pPr.insert_element_before(el, *PPR_ORDER[PPR_ORDER.index(tag) + 1:])
    for k, v in (attrs or {}).items():
        el.set(qn(k), str(v))
    return el


def set_run_fonts(rPr, font):
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.insert(0, rFonts)
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        rFonts.attrib.pop(qn(attr), None)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rFonts.set(qn(attr), font)


def set_lang(rPr, lang):
    el = rPr.find(qn("w:lang"))
    if el is None:
        el = OxmlElement("w:lang")
        rPr.insert_element_before(el, "w:eastAsianLayout", "w:specVanish", "w:oMath")
    el.set(qn("w:val"), lang)
    el.set(qn("w:eastAsia"), lang)


def add_field(par, instr, placeholder=""):
    """Вставляет поле Word (PAGE, TOC, SEQ ...) с текстом-заглушкой до обновления."""
    def fld_char(kind):
        r = par.add_run()
        fc = OxmlElement("w:fldChar")
        fc.set(qn("w:fldCharType"), kind)
        r._r.append(fc)
        return r
    fld_char("begin")
    r = par.add_run()
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = f" {instr} "
    r._r.append(it)
    fld_char("separate")
    result = par.add_run(placeholder)
    fld_char("end")
    return result


def add_abstract_num(doc, levels_xml, multilevel=False):
    """Создаёт описание списка в numbering.xml и возвращает его numId."""
    numbering = doc.part.numbering_part.element
    abs_id = max([int(a.get(qn("w:abstractNumId"))) for a in numbering.findall(qn("w:abstractNum"))] + [0]) + 1
    num_id = max([int(n.get(qn("w:numId"))) for n in numbering.findall(qn("w:num"))] + [0]) + 1
    kind = "multilevel" if multilevel else "singleLevel"
    abstract = parse_xml(
        f'<w:abstractNum {nsdecls("w")} w:abstractNumId="{abs_id}">'
        f'<w:multiLevelType w:val="{kind}"/>{levels_xml}</w:abstractNum>'
    )
    first_num = numbering.find(qn("w:num"))
    if first_num is not None:
        first_num.addprevious(abstract)
    else:
        numbering.append(abstract)
    numbering.append(parse_xml(
        f'<w:num {nsdecls("w")} w:numId="{num_id}"><w:abstractNumId w:val="{abs_id}"/></w:num>'
    ))
    return num_id


def lvl_xml(ilvl, fmt, text, indent_twips, font, style_id=None, bold=False):
    pstyle = f'<w:pStyle w:val="{style_id}"/>' if style_id else ""
    b = "<w:b/><w:bCs/>" if bold else ""
    return (
        f'<w:lvl w:ilvl="{ilvl}"><w:start w:val="1"/><w:numFmt w:val="{fmt}"/>{pstyle}'
        f'<w:suff w:val="space"/><w:lvlText w:val="{text}"/><w:lvlJc w:val="left"/>'
        f'<w:pPr><w:ind w:left="0" w:firstLine="{indent_twips}"/></w:pPr>'
        f'<w:rPr><w:rFonts w:ascii="{font}" w:hAnsi="{font}" w:cs="{font}"/>{b}</w:rPr></w:lvl>'
    )


def link_numbering(style, num_id, ilvl=0):
    numPr = style.element.get_or_add_pPr().get_or_add_numPr()
    numPr.get_or_add_ilvl().val = ilvl
    numPr.get_or_add_numId().val = num_id


# ---------------------------------------------------------------- стили

class Builder:
    def __init__(self, cfg):
        self.cfg = cfg
        self.doc = Document()
        self.t = cfg["text"]
        self.font = self.t["font"]
        page = cfg["page"]
        m = page["margins_mm"]
        self.text_width_twips = int((page["width_mm"] - m["left"] - m["right"]) / 10 * TWIPS_PER_CM)
        self.indent_twips = int(self.t["first_line_indent_cm"] * TWIPS_PER_CM)
        fields = dict(cfg.get("fields", {}))
        fields["year"] = fields.get("year") or str(datetime.date.today().year)
        self.fields = Fields({k: v for k, v in fields.items() if not k.startswith("_")})

    def fill(self, s):
        return s.format_map(self.fields)

    def style(self, name, base="Normal", kind=WD_STYLE_TYPE.PARAGRAPH):
        styles = self.doc.styles
        try:
            st = styles[name]
        except KeyError:
            st = styles.add_style(name, kind)
            st.base_style = styles[base]
            st.quick_style = True
        return st

    def fmt(self, st, *, size=None, bold=False, italic=False, align="left", first_line_cm=0.0,
            left_cm=0.0, line_spacing=None, before=0, after=0, page_break=False, keep_next=False,
            no_hyphens=False, outline=None):
        f = st.font
        f.name = self.font
        f.size = Pt(size or self.t["size_pt"])
        f.bold = bold
        f.italic = italic
        f.color.rgb = RGBColor(0, 0, 0)
        rPr = st.element.get_or_add_rPr()
        for attr in ("w:themeColor", "w:themeShade", "w:themeTint"):  # у встроенных заголовков синий цвет темы
            rPr.find(qn("w:color")).attrib.pop(qn(attr), None)
        set_run_fonts(rPr, self.font)
        set_lang(rPr, self.t.get("language", "ru-RU"))
        pf = st.paragraph_format
        pf.alignment = ALIGN[align]
        pf.first_line_indent = Cm(first_line_cm)
        pf.left_indent = Cm(left_cm)
        pf.right_indent = Cm(0)
        pf.line_spacing = line_spacing if line_spacing is not None else self.t["line_spacing"]
        pf.space_before = Pt(before)
        pf.space_after = Pt(after)
        pf.page_break_before = page_break
        pf.keep_with_next = keep_next
        pf.widow_control = True
        pPr = st.element.get_or_add_pPr()
        if no_hyphens:
            ppr_set(pPr, "w:suppressAutoHyphens")
        ppr_set(pPr, "w:contextualSpacing", {"w:val": "0"})
        if outline is not None:
            ppr_set(pPr, "w:outlineLvl", {"w:val": outline})
        return st

    def heading_fmt(self, st, h, outline):
        centered = h["align"] == "center"
        return self.fmt(
            st, size=h["size_pt"], bold=h["bold"], align="center" if centered else "left",
            first_line_cm=0 if centered else self.t["first_line_indent_cm"],
            before=h["space_before_pt"], after=h["space_after_pt"], page_break=h["new_page"],
            keep_next=True, no_hyphens=True, outline=outline,
        )

    def build_styles(self):
        t, cap, hd = self.t, self.cfg["captions"], self.cfg["headings"]
        self.fmt(self.doc.styles["Normal"], align=t["align"], first_line_cm=t["first_line_indent_cm"])

        # Заголовки глав с автоматической нумерацией 1 / 1.1 / 1.1.1 (без точки в конце, ГОСТ 7.32)
        heading_styles = [self.doc.styles[f"Heading {i}"] for i in (1, 2, 3)]
        texts = ["%1", "%1.%2", "%1.%2.%3"]
        levels = "".join(
            lvl_xml(i, "decimal", texts[i], 0 if hd[f"level{i+1}"]["align"] == "center" else self.indent_twips,
                    self.font, heading_styles[i].style_id, hd[f"level{i+1}"]["bold"])
            for i in range(3)
        )
        num_id = add_abstract_num(self.doc, levels, multilevel=True)
        for i, st in enumerate(heading_styles):
            self.heading_fmt(st, hd[f"level{i+1}"], outline=i)
            link_numbering(st, num_id, i)

        # Ненумеруемые заголовки: ВВЕДЕНИЕ, ЗАКЛЮЧЕНИЕ, ПРИЛОЖЕНИЕ... (попадают в оглавление)
        self.heading_fmt(self.style("Структурный заголовок"), hd["structural"], outline=0)
        # «СОДЕРЖАНИЕ» — так же выглядит, но в оглавление не попадает
        self.heading_fmt(self.style("Заголовок содержания"), hd["structural"], outline=9)
        self.fmt(self.style("Название приложения"), bold=True, align="center", after=12, keep_next=True,
                 no_hyphens=True)

        for i in (1, 2, 3):
            self.fmt(self.style(f"toc {i}"), align="left", left_cm=0.5 * (i - 1),
                     line_spacing=t["line_spacing"])

        self.fmt(self.style("Рисунок"), align="center", keep_next=True, line_spacing=1.0, before=6, after=6)
        self.fmt(self.style("Подпись рисунка"), size=cap["figure"]["size_pt"], align=cap["figure"]["align"],
                 after=cap["figure"]["space_after_pt"], no_hyphens=True)
        self.fmt(self.style("Подпись таблицы"), size=cap["table"]["size_pt"], align=cap["table"]["align"],
                 keep_next=True, before=6, no_hyphens=True)
        self.fmt(self.style("Текст таблицы"), size=cap["table_text_size_pt"], align="left", line_spacing=1.0)
        self.fmt(self.style("Формула"), align="left", before=6, after=6)
        tabs = self.style("Формула").paragraph_format.tab_stops
        tabs.add_tab_stop(Cm(self.text_width_twips / TWIPS_PER_CM / 2), WD_TAB_ALIGNMENT.CENTER)
        tabs.add_tab_stop(Cm(self.text_width_twips / TWIPS_PER_CM), WD_TAB_ALIGNMENT.RIGHT)
        self.fmt(self.style("Пояснение к формуле"), align=t["align"],
                 first_line_cm=t["first_line_indent_cm"] if cap["formula_where_indent"] else 0)
        self.fmt(self.style("Титульный лист"), size=self.cfg["title_page"]["font_size_pt"], align="center",
                 line_spacing=1.0, no_hyphens=True)

        # Списки: маркированный «–» и нумерованные «1)» / «1.» — с абзацного отступа
        bullet = self.cfg["lists"]["bullet"]
        for name, fmt_, txt in (("Список тире", "bullet", bullet), ("Список нумерованный", "decimal", "%1)"),
                                ("Список источников", "decimal", "%1.")):
            st = self.fmt(self.style(name), align=t["align"], first_line_cm=t["first_line_indent_cm"])
            link_numbering(st, add_abstract_num(self.doc, lvl_xml(0, fmt_, txt, self.indent_twips, self.font)))

    # ---------------------------------------------------------------- страница, титул

    def setup_page(self):
        page = self.cfg["page"]
        sec = self.doc.sections[0]
        sec.page_width, sec.page_height = Mm(page["width_mm"]), Mm(page["height_mm"])
        m = page["margins_mm"]
        sec.left_margin, sec.right_margin = Mm(m["left"]), Mm(m["right"])
        sec.top_margin, sec.bottom_margin = Mm(m["top"]), Mm(m["bottom"])
        sec.footer_distance = Mm(10)

        num = page["numbering"]
        p = sec.footer.paragraphs[0]
        self._plain(p, align=num["align"])
        run = add_field(p, "PAGE", "2")
        run.font.size = Pt(num["size_pt"])

        if not num["show_on_title"]:
            sec.different_first_page_header_footer = True
            first = sec.first_page_footer.paragraphs[0]
            self._plain(first, align="center")
            first.style = self.doc.styles["Титульный лист"]
            first.text = "\n".join(self.fill(x) for x in self.cfg["title_page"].get("bottom_lines", []))

        settings = self.doc.settings.element
        if self.t.get("hyphenation"):
            el = OxmlElement("w:autoHyphenation")
            tab = settings.find(qn("w:defaultTabStop"))
            tab.addnext(el) if tab is not None else settings.insert(0, el)

    def _plain(self, p, align="left"):
        pf = p.paragraph_format
        pf.first_line_indent = Cm(0)
        pf.line_spacing = 1.0
        pf.alignment = ALIGN[align]

    def title_page(self):
        tp = self.cfg["title_page"]
        style = self.doc.styles["Титульный лист"]
        for block in tp["blocks"]:
            first = True
            if block["type"] == "signatures":
                self._signature_table(block)
                continue
            for line in block["lines"]:
                p = self.doc.add_paragraph(style=style)
                p.alignment = ALIGN[block.get("align", "center")]
                if block.get("indent_left_cm"):
                    p.paragraph_format.left_indent = Cm(block["indent_left_cm"])
                if first and block.get("space_before_pt"):
                    p.paragraph_format.space_before = Pt(block["space_before_pt"])
                first = False
                run = p.add_run(self.fill(line))
                run.bold = block.get("bold", False)
                if block.get("size_pt"):
                    run.font.size = Pt(block["size_pt"])

    def _signature_table(self, block):
        spacer = self.doc.add_paragraph(style="Титульный лист")
        spacer.paragraph_format.space_before = Pt(block.get("space_before_pt", 0))
        widths = [Cm(w) for w in block["col_widths_cm"]]
        table = self.doc.add_table(rows=0, cols=len(widths))
        table.autofit = False
        for row in block["rows"]:
            cells = table.add_row().cells
            for i, text in enumerate(row):
                cells[i].width = widths[i]
                lines = self.fill(text).split("\n")
                cells[i].paragraphs[0].text = lines[0]
                for extra in lines[1:]:
                    cells[i].add_paragraph(extra)
                paras = cells[i].paragraphs
                for n, p in enumerate(paras):
                    p.style = self.doc.styles["Титульный лист"]
                    p.alignment = ALIGN["left"]
                    p.paragraph_format.space_after = Pt(12) if n == len(paras) - 1 else Pt(0)
                if i == len(row) - 1:  # подпись прижимаем к низу ячейки
                    cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.BOTTOM

    # ---------------------------------------------------------------- основная часть

    def case(self, text, h):
        return text.upper() if h.get("case") == "upper" else text

    def body(self, items):
        for item in items or []:
            if isinstance(item, str):
                self.doc.add_paragraph(item)
            elif "dash" in item:
                for x in item["dash"]:
                    self.doc.add_paragraph(x, style="Список тире")
            elif "num" in item:
                for x in item["num"]:
                    self.doc.add_paragraph(x, style="Список нумерованный")
            elif "h3" in item:
                self.doc.add_paragraph(item["h3"], style="Heading 3")
            elif "figure" in item:
                self.figure(item["figure"])
            elif "table" in item:
                self.table(item["table"])
            elif "formula" in item:
                self.formula(item["formula"])

    def caption(self, style, prefix, seq, title):
        p = self.doc.add_paragraph(style=style)
        p.add_run(prefix + " ")
        add_field(p, f"SEQ {seq} \\* ARABIC", "1")
        p.add_run(self.cfg["captions"]["separator"] + title)

    def figure(self, title):
        p = self.doc.add_paragraph("[Место для рисунка]", style="Рисунок")
        pPr = p._p.get_or_add_pPr()
        bdr = ppr_set(pPr, "w:pBdr")
        for side in ("top", "left", "bottom", "right"):
            bdr.append(parse_xml(f'<w:{side} {nsdecls("w")} w:val="dashed" w:sz="6" w:space="24" w:color="808080"/>'))
        p.paragraph_format.left_indent = Cm(4)
        p.paragraph_format.right_indent = Cm(4)
        p.paragraph_format.space_before = Pt(30)
        p.paragraph_format.space_after = Pt(30)
        self.caption("Подпись рисунка", self.cfg["captions"]["figure"]["prefix"], "Рисунок", title)

    def table(self, spec):
        self.caption("Подпись таблицы", self.cfg["captions"]["table"]["prefix"], "Таблица", spec["caption"])
        ncols = len(spec["header"])
        table = self.doc.add_table(rows=1, cols=ncols)
        table.style = self.doc.styles["Table Grid"]
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        col_w = Cm(self.text_width_twips / TWIPS_PER_CM / ncols)
        hdr = table.rows[0]
        trPr = hdr._tr.get_or_add_trPr()
        trPr.append(OxmlElement("w:tblHeader"))  # повтор шапки на каждой странице
        for i, text in enumerate(spec["header"]):
            cell = hdr.cells[i]
            cell.width = col_w
            p = cell.paragraphs[0]
            p.style = self.doc.styles["Текст таблицы"]
            p.alignment = ALIGN["center"]
            p.add_run(text).bold = True
        for row in spec["rows"]:
            cells = table.add_row().cells
            for i, text in enumerate(row):
                cells[i].width = col_w
                cells[i].paragraphs[0].style = self.doc.styles["Текст таблицы"]
                cells[i].paragraphs[0].text = text
        self.doc.add_paragraph()  # отбивка после таблицы

    def formula(self, spec):
        p = self.doc.add_paragraph(style="Формула")
        p.add_run("\t")
        p.add_run(spec["expr"]).italic = True
        p.add_run("\t(")
        add_field(p, "SEQ Формула \\* ARABIC", "1")
        p.add_run(")")
        for line in spec.get("where", []):
            self.doc.add_paragraph(line, style="Пояснение к формуле")

    def toc(self, title):
        hd = self.cfg["headings"]
        self.doc.add_paragraph(self.case(title, hd["structural"]), style="Заголовок содержания")
        p = self.doc.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(0)
        add_field(p, f'TOC \\o "1-{hd["toc_levels"]}" \\h \\z \\u',
                  "Оглавление обновится при открытии в Word (или: правой кнопкой → «Обновить поле»).")

    def structure(self):
        hd = self.cfg["headings"]
        for item in self.cfg["structure"]:
            kind = item["type"]
            if kind == "toc":
                self.toc(item["title"])
            elif kind == "structural":
                self.doc.add_paragraph(self.case(item["title"], hd["structural"]), style="Структурный заголовок")
                self.body(item.get("body"))
            elif kind == "chapter":
                self.doc.add_paragraph(self.case(item["title"], hd["level1"]), style="Heading 1")
                for sec in item.get("sections", []):
                    self.doc.add_paragraph(self.case(sec["title"], hd["level2"]), style="Heading 2")
                    self.body(sec.get("body"))
            elif kind == "references":
                self.doc.add_paragraph(self.case(item["title"], hd["structural"]), style="Структурный заголовок")
                for x in item["items"]:
                    self.doc.add_paragraph(x, style="Список источников")
            elif kind == "appendix":
                self.doc.add_paragraph(self.case(f"Приложение {item['letter']}", hd["structural"]),
                                       style="Структурный заголовок")
                self.doc.add_paragraph(item["title"], style="Название приложения")
                self.body(item.get("body"))

    def build(self):
        self.build_styles()
        self.setup_page()
        self.title_page()
        self.structure()
        self.doc.core_properties.title = self.fill(self.fields.get("topic", ""))
        self.doc.core_properties.author = self.fill(self.fields.get("student", ""))
        return self.doc


# ---------------------------------------------------------------- обновление полей через Word

WORD_PS = r"""
$ErrorActionPreference = 'Stop'
$w = New-Object -ComObject Word.Application
$w.Visible = $false; $w.DisplayAlerts = 0
try {
  $d = $w.Documents.Open('{docx}')
  if ($null -eq $d) { throw 'Word не открыл документ' }
  $d.Fields.Update() | Out-Null
  foreach ($t in $d.TablesOfContents) { $t.Update() }
  $d.Save()
  if ('{pdf}') { $d.ExportAsFixedFormat('{pdf}', 17) }
  $d.Close()
} finally { $w.Quit() }
"""


def update_with_word(docx: Path, pdf: Path | None) -> bool:
    """Обновляет поля через Word. Работает с копией в короткой временной папке:
    Word не открывает файлы с путём длиннее ~255 символов."""
    if sys.platform != "win32" or not shutil.which("powershell"):
        return False
    with tempfile.TemporaryDirectory() as tmp:
        work_docx, work_pdf = Path(tmp) / "t.docx", Path(tmp) / "t.pdf"
        shutil.copyfile(docx, work_docx)
        script = WORD_PS.replace("{docx}", str(work_docx)).replace("{pdf}", str(work_pdf) if pdf else "")
        res = subprocess.run(["powershell", "-NoProfile", "-NonInteractive", "-Command", script],
                             capture_output=True, text=True, stdin=subprocess.DEVNULL,
                             creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
        if res.returncode != 0:
            return False
        shutil.copyfile(work_docx, docx)
        if pdf and work_pdf.exists():
            shutil.copyfile(work_pdf, pdf)
    return True


def mark_fields_dirty(docx: Path):
    """Без Word: просим Word обновить оглавление при первом открытии."""
    doc = Document(docx)
    el = OxmlElement("w:updateFields")
    el.set(qn("w:val"), "true")
    doc.settings.element.insert_element_before(
        el, "w:hdrShapeDefaults", "w:footnotePr", "w:endnotePr", "w:compat", "w:docVars", "w:rsids")
    doc.save(docx)


def generate(cfg: dict, out: Path, pdf: bool = False, use_word: bool = True):
    """Собирает шаблон. Возвращает (путь к docx, путь к pdf или None, обновлял ли Word поля)."""
    out = Path(out)
    out.parent.mkdir(parents=True, exist_ok=True)
    Builder(cfg).build().save(out)
    pdf_path = out.with_suffix(".pdf") if pdf else None
    word = use_word and update_with_word(out, pdf_path)
    if not word:
        mark_fields_dirty(out)
    return out, (pdf_path if pdf_path and pdf_path.exists() else None), word


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("config", type=Path, help="JSON с правилами оформления")
    ap.add_argument("-o", "--out", type=Path, help="куда сохранить .docx")
    ap.add_argument("--pdf", action="store_true", help="дополнительно сохранить PDF (нужен Word)")
    ap.add_argument("--no-word", action="store_true", help="не запускать Word для обновления полей")
    args = ap.parse_args()

    cfg = load_config(args.config)
    out = args.out or Path("out") / f"{cfg['name']}.docx"
    docx, pdf, word = generate(cfg, out, pdf=args.pdf, use_word=not args.no_word)
    print("Готово:", docx, f"+ {pdf}" if pdf else "", "" if word else "(поля обновятся при открытии в Word)")


if __name__ == "__main__":
    main()
